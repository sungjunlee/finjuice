"""Actual migration exports, deterministic artifacts and stale receipt checks."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
from typer.testing import CliRunner

from finjuice.pipeline.cli.main import app
from finjuice.pipeline.export import source
from finjuice.pipeline.storage.mutation_facade import ConfigDocument, StorageMutationFacade
from finjuice.pipeline.storage.read_facade import read_transaction_snapshot
from tests.cli.commands.test_repository_mutation_fences import active_root as _empty_fixture
from tests.cli.commands.test_repository_query import QueryRoot
from tests.cli.commands.test_repository_query import query_root as _query_fixture

query_root = _query_fixture
empty_root = _empty_fixture


def _export(root: QueryRoot, *args: str, legacy: bool = False, no_filter: bool = False):
    argv = ["--data-dir", str(root.legacy if legacy else root.root)]
    if no_filter:
        argv.append("--no-filter")
    return CliRunner().invoke(
        app,
        [*argv, "export", "--no-auto-open", "--json", *args],
        obj={} if legacy else {"activation_evidence_provider": root.provider},
    )


def _verify(root: QueryRoot, manifest: Path, *, human: bool = False):
    return CliRunner().invoke(
        app,
        [
            "--data-dir",
            str(root.root),
            "export-verify",
            str(manifest),
            *([] if human else ["--json"]),
        ],
        obj={"activation_evidence_provider": root.provider},
    )


def _payload(result):
    assert result.exit_code == 0, result.output
    return json.loads(result.output)


def _files(payload):
    return {item["kind"]: Path(item["path"]) for item in payload["output_files"]}


def _replace_rules(root: QueryRoot, content: bytes):
    facade = StorageMutationFacade(root.root, root.provider)
    return facade.replace_config(
        ConfigDocument.from_validated_yaml("rules", content, parser_version="test")
    )


def test_export_preserves_baseline_master_and_filtered_reports(query_root: QueryRoot):
    before = read_transaction_snapshot(query_root.root, query_root.provider)
    old = _payload(_export(query_root, legacy=True))
    new = _payload(_export(query_root))
    assert new["transaction_count"] == old["transaction_count"] == 2
    assert new["_meta"]["filters_applied"] == 1
    assert new["_meta"]["dataset_generation"] == query_root.generation
    assert new["_meta"]["dataset_revision"] == 0
    assert new["_meta"]["calculation_as_of"] == "2026-09-01"
    old_files, new_files = _files(old), _files(new)
    assert new_files.keys() - {"transactions_csv"} == old_files.keys()
    assert new_files["transactions_csv"].exists()
    for kind in old_files:
        if kind == "master_xlsx":
            books = [
                openpyxl.load_workbook(p, read_only=True)
                for p in (old_files[kind], new_files[kind])
            ]
            try:
                rows = [list(book.active.values) for book in books]
                header = rows[1][0]
                assert header[: len(rows[0][0])] == rows[0][0]
                legacy_rows = [list(row) for row in rows[0][1:]]
                manual_index = header.index("tags_manual")
                for row in legacy_rows:
                    row[manual_index] = "visible"
                assert legacy_rows == [list(row[: len(rows[0][0])]) for row in rows[1][1:]]
                assert rows[1][1][header.index("category_manual")] == "manual"
                assert rows[1][1][header.index("transaction_id")] != rows[1][1][0]
                assert rows[1][1][header.index("amount_coefficient")] == "'-120025"
                assert rows[1][1][header.index("notes_manual")] == "manual\nnotes"
                assert rows[1][1][header.index("tags_final")] == "persisted, persisted"
            finally:
                for book in books:
                    book.close()
        else:
            assert new_files[kind].read_bytes() == old_files[kind].read_bytes(), kind
    manifest = Path(new["manifest_path"])
    verification = _payload(_verify(query_root, manifest))
    assert verification["stale"] is False
    assert verification["integrity"] == "intact"
    from tests.test_json_schemas import _load_schema, _validator_for

    _validator_for(_load_schema("export_verify.schema.json")).validate(verification)
    assert _verify(query_root, manifest, human=True).exit_code == 0
    assert read_transaction_snapshot(query_root.root, query_root.provider) == before


def test_export_all_bytes_replay_and_ignores_live_csv_rules(query_root: QueryRoot, monkeypatch):
    first = _payload(_export(query_root, "--format", "all", "--period", "2026-09"))
    (query_root.root / "rules.yaml").write_text("rules: [invalid")
    csv = query_root.root / "transactions/2026/09/transactions.csv"
    csv.parent.mkdir(parents=True)
    csv.write_text("amount\n999999\n")
    from datetime import datetime

    from finjuice.pipeline.export import result as export_result

    class LaterClock:
        @classmethod
        def now(cls):
            return datetime(2035, 2, 1)

    monkeypatch.setattr(export_result, "datetime", LaterClock)
    second = _payload(_export(query_root, "--format", "all", "--period", "2026-09"))
    a, b = _files(first), _files(second)
    assert a.keys() == b.keys()
    assert {"html_report", "markdown_report", "master_xlsx"} <= a.keys()
    for kind in a:
        assert a[kind].read_bytes() == b[kind].read_bytes(), kind
    assert Path(first["manifest_path"]).read_bytes() == Path(second["manifest_path"]).read_bytes()
    assert first["manifest_path"] != second["manifest_path"]


def test_export_rules_mutation_during_render_stays_pinned_then_becomes_stale(
    query_root: QueryRoot, monkeypatch
):
    original = source.read_transaction_snapshot
    calls = []

    def snapshot_then_mutate(*args, **kwargs):
        snapshot = original(*args, **kwargs)
        calls.append(snapshot.info.dataset_revision)
        _replace_rules(query_root, b"rules: []\nreport_filters: {}\n")
        return snapshot

    monkeypatch.setattr(source, "read_transaction_snapshot", snapshot_then_mutate)
    result = _payload(_export(query_root))
    assert calls == [0]
    assert result["_meta"]["dataset_revision"] == 0
    assert result["_meta"]["filters_applied"] == 1
    check = _payload(_verify(query_root, Path(result["manifest_path"])))
    assert check["stale"] is True
    assert check["current"]["dataset_revision"] == 1
    assert check["integrity"] == "intact"


def test_export_detects_modified_and_missing_artifacts_separately(query_root: QueryRoot):
    result = _payload(_export(query_root))
    files = _files(result)
    paths = [files["transactions_csv"], files["master_xlsx"]]
    paths[0].write_bytes(b"changed derived file")
    paths[1].unlink()
    check = _payload(_verify(query_root, Path(result["manifest_path"])))
    assert check["stale"] is False
    assert check["integrity"] == "mismatch"
    assert {"modified", "missing"} <= {item["status"] for item in check["files"]}
    assert (
        read_transaction_snapshot(query_root.root, query_root.provider).info.dataset_revision == 0
    )


def test_export_dry_run_reads_snapshot_counts_without_writing(query_root: QueryRoot):
    result = _payload(_export(query_root, "--dry-run", "--format", "md", "--period", "2026-08"))
    assert result["transaction_count"] == 0
    assert result["_meta"]["dataset_revision"] == 0
    assert not (query_root.root / "exports").exists()
    unfiltered = _payload(_export(query_root, "--dry-run", no_filter=True))
    assert unfiltered["transaction_count"] == 2
    assert unfiltered["_meta"]["filters_applied"] == 0
    assert not (query_root.root / "exports").exists()


def test_export_all_filtered_reports_never_claim_old_files(query_root: QueryRoot):
    _replace_rules(
        query_root,
        b"rules: []\nreport_filters:\n  excluded_categories:\n"
        b"    - {name: kept, reason: test}\n    - {name: excluded, reason: test}\n",
    )
    old = query_root.root / "exports/reports/by_category.csv"
    old.parent.mkdir(parents=True)
    old.write_bytes(b"OLD PRIVATE DERIVED FILE")
    result = _payload(_export(query_root))
    assert set(_files(result)) == {"master_xlsx", "transactions_csv"}
    assert old.read_bytes() == b"OLD PRIVATE DERIVED FILE"
    assert _payload(_verify(query_root, Path(result["manifest_path"])))["integrity"] == "intact"


def test_export_without_evidence_cannot_fall_back(query_root: QueryRoot):
    result = CliRunner().invoke(app, ["--data-dir", str(query_root.root), "export", "--json"])
    assert result.exit_code != 0
    assert not (query_root.root / "exports").exists()


@pytest.mark.parametrize("invalid_path", ["../outside", "/tmp/outside", "reports/../../outside"])
def test_export_manifest_rejects_escaping_paths(query_root: QueryRoot, invalid_path: str):
    result = _payload(_export(query_root))
    path = Path(result["manifest_path"])
    manifest = json.loads(path.read_text())
    manifest["files"][0]["path"] = invalid_path
    path.write_text(json.dumps(manifest))
    assert _verify(query_root, path).exit_code != 0


@pytest.mark.parametrize("status", ["invalid", "opaque"])
def test_export_invalid_rules_status_only_allows_explicit_filter_bypass(
    query_root: QueryRoot, status: str
):
    marker = "SYNTHETIC_PRIVATE_CONFIG"
    StorageMutationFacade(query_root.root, query_root.provider).replace_config(
        ConfigDocument("rules", f"rules: []\n# {marker}\n".encode(), status, None, "test")
    )
    result = _export(query_root)
    assert result.exit_code != 0
    assert marker not in result.output
    assert not (query_root.root / "exports").exists()
    bypass = _payload(_export(query_root, no_filter=True))
    assert bypass["_meta"]["filters_disabled"] is True


def test_export_empty_repository_does_not_reuse_existing_master(empty_root):
    root = QueryRoot(empty_root.root, empty_root.root, empty_root.provider, empty_root.generation)
    old = root.root / "exports/master_undated.xlsx"
    old.write_bytes(b"OLD")
    result = _payload(_export(root))
    assert result["transaction_count"] == 0
    assert set(_files(result)) == {"transactions_csv"}
    assert old.read_bytes() == b"OLD"
    assert result["_meta"]["calculation_as_of"] is None
    check = _payload(_verify(root, Path(result["manifest_path"])))
    assert check["stale"] is False
    assert len(check["files"]) == 1
    assert check["files"][0]["path"] == "transactions.csv"


def test_export_manifest_rejects_symlink_member(query_root: QueryRoot, tmp_path: Path):
    result = _payload(_export(query_root))
    path = next(iter(_files(result).values()))
    path.unlink()
    other = tmp_path / "outside"
    other.write_bytes(b"PRIVATE")
    path.symlink_to(other)
    verification = _verify(query_root, Path(result["manifest_path"]))
    assert verification.exit_code != 0
    assert "PRIVATE" not in verification.output


def test_export_failure_discards_staging_and_keeps_previous_run(query_root, monkeypatch):
    from finjuice.pipeline.export import result as export_result

    previous = _payload(_export(query_root))
    previous_manifest = Path(previous["manifest_path"])
    before = previous_manifest.read_bytes()

    def fail_after_file(run):
        (run.paths.export_dir / "partial.xlsx").write_bytes(b"PARTIAL")
        raise RuntimeError("SYNTHETIC_PRIVATE_ERROR")

    monkeypatch.setattr(export_result, "_generate_xlsx_outputs", fail_after_file)
    failure = _export(query_root)
    assert failure.exit_code != 0
    assert "SYNTHETIC_PRIVATE_ERROR" not in failure.output
    assert previous_manifest.read_bytes() == before
    assert len(list((query_root.root / "exports/runs").iterdir())) == 1
    assert not list((query_root.root / "exports").glob(".repository-export-*"))


def test_export_period_does_not_filter_master_in_all_format(query_root: QueryRoot):
    result = _payload(_export(query_root, "--format", "all", "--period", "2026-08"))
    assert result["transaction_count"] == 2
    assert _files(result)["master_xlsx"].exists()
    markdown = _files(result)["markdown_report"].read_text()
    assert "shop-0" not in markdown and "shop-1" not in markdown


def test_export_native_null_aliases_keep_uuid_and_exact_audit_fields(tmp_path: Path):
    import shutil

    from finjuice.pipeline.storage.authority import AuthorityPaths, StaticActivationEvidenceProvider
    from finjuice.pipeline.storage.sqlite import RepositoryReader
    from tests.pipeline.test_sqlite_bulk_mutations import _evidence, _write_activation
    from tests.pipeline.test_sqlite_transaction_scopes import _native

    candidate = _native(tmp_path, ["2026-09-01", None])
    with RepositoryReader(candidate.database) as reader:
        generation = reader.info.dataset_generation
    root = tmp_path / "active"
    paths = AuthorityPaths.for_data_dir(root)
    shutil.copytree(candidate.root, paths.generation(generation).root)
    _write_activation(paths, generation)
    case = QueryRoot(
        root, tmp_path / "unused", StaticActivationEvidenceProvider(_evidence()), generation
    )
    result = _payload(_export(case))
    workbook = openpyxl.load_workbook(_files(result)["master_xlsx"], read_only=True)
    try:
        rows = list(workbook.active.values)
        header, data = rows[0], rows[1:]
        assert len(data) == 2
        assert all(row[header.index("row_hash")] is None for row in data)
        ids = [row[header.index("transaction_id")] for row in data]
        assert len(set(ids)) == 2 and all(ids)
        assert "amount_coefficient" in header and "amount_lexical" in header
    finally:
        workbook.close()


def test_export_all_does_not_claim_partially_written_dependency_failure(query_root, monkeypatch):
    from finjuice.pipeline.export import html_report

    def partial(options):
        options.output_path.write_text("PARTIAL_PRIVATE")
        raise ImportError("PRIVATE_DEPENDENCY")

    monkeypatch.setattr(html_report, "render_html_report", partial)
    result = _export(query_root, "--format", "all")
    assert result.exit_code != 0
    assert "PRIVATE" not in result.output
    assert not list((query_root.root / "exports").glob(".repository-export-*"))
    assert not (query_root.root / "exports/runs").exists()


def test_export_dry_run_paths_describe_new_run(query_root: QueryRoot):
    result = _payload(_export(query_root, "--dry-run"))
    assert all("/runs/<new-run>/" in item["path"] for item in result["output_files"])
    assert all(item["would_overwrite"] is False for item in result["output_files"])
    assert any(item["path"].endswith("master_20260901.xlsx") for item in result["output_files"])


def test_export_report_failure_is_not_misreported_as_zero_rows(query_root, monkeypatch):
    from finjuice.pipeline.export import reports

    def fail_report(*args):
        raise RuntimeError("SYNTHETIC_PRIVATE_AGGREGATION")

    monkeypatch.setitem(reports._REPORT_EXPORTERS, "by_category", fail_report)
    result = _export(query_root)
    assert result.exit_code != 0
    assert "SYNTHETIC_PRIVATE_AGGREGATION" not in result.output
    assert not (query_root.root / "exports/runs").exists()


def test_export_row_mutation_during_render_keeps_original_values(query_root, monkeypatch):
    import csv

    from finjuice.pipeline.export import result as export_result
    from finjuice.pipeline.storage.sqlite.mutations import ManualTransactionEdit

    original = export_result._generate_xlsx_outputs

    def mutate_then_render(run):
        identifier = run.full_source_df["transaction_id"][0]
        StorageMutationFacade(query_root.root, query_root.provider).edit_manual_transaction(
            ManualTransactionEdit(identifier=identifier, note_supplied=True, note="later note")
        )
        return original(run)

    monkeypatch.setattr(export_result, "_generate_xlsx_outputs", mutate_then_render)
    result = _payload(_export(query_root))
    with _files(result)["transactions_csv"].open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert {row["notes_manual"] for row in rows} == {"manual\nnotes"}
    check = _payload(_verify(query_root, Path(result["manifest_path"])))
    assert check["source"]["dataset_revision"] == 0
    assert check["current"]["dataset_revision"] == 1
    assert check["stale"] is True and check["integrity"] == "intact"
